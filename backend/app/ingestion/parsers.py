"""
Parsing layer: raw bytes -> ParsedDocument (structured text, tables,
pictures, with page numbers and section headings for citations).

Each parser is single-purpose; the pipeline picks the parser by
content_type. Adding a new format means adding one function here and one
line in the dispatch table.

Phase 2 (multimodal upgrade): PDF and DOCX now go through Docling, which
returns typed elements (text, headings, lists, tables, pictures) instead
of a flat string. The pipeline for Phase 2 consumes **text_units only** —
table_units and picture_units are captured but not yet routed to chunks
(that's Phase 3 for tables, Phase 4 for pictures).

TXT/MD keep the plain decoder — routing simple text through a layout
model adds cost with no benefit.
"""

from __future__ import annotations

import logging
import os
import tempfile
from dataclasses import dataclass, field
from functools import lru_cache
from io import BytesIO

from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions
from docling.document_converter import DocumentConverter, PdfFormatOption

from app.core.exceptions import IngestionError

logger = logging.getLogger("rag.ingestion.parsers")


# ------------------------------------------------------------------ types

@dataclass
class ParsedUnit:
    """A run of text with optional page/section metadata."""
    text: str
    page_number: int | None = None
    section_title: str | None = None


@dataclass
class TableUnit:
    """A table serialized to markdown (what the LLM will see) plus
    optional future structured data (rows/cells)."""
    markdown: str
    page_number: int | None = None
    section_title: str | None = None


@dataclass
class PictureUnit:
    """An image extracted from the document, PNG-encoded. Phase 4 adds
    the vision-model description; Phase 2 only captures the bytes."""
    image_bytes: bytes
    page_number: int | None = None
    section_title: str | None = None
    caption: str | None = None


@dataclass
class ParsedDocument:
    """Structured parse output. Phase 2 consumes text_units only; Phase 3
    consumes table_units; Phase 4 consumes picture_units."""
    text_units: list[ParsedUnit] = field(default_factory=list)
    table_units: list[TableUnit] = field(default_factory=list)
    picture_units: list[PictureUnit] = field(default_factory=list)


# ------------------------------------------------------------------ Docling

@lru_cache(maxsize=1)
def _get_docling_converter() -> DocumentConverter:
    """One converter per process — instantiation loads the layout model
    (30–90 s first time, cached afterward). Picture image generation is
    enabled so Phase 4 can pick up PictureUnits without re-parsing."""
    pipeline_options = PdfPipelineOptions()
    pipeline_options.generate_picture_images = True
    pipeline_options.images_scale = 2.0

    return DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options),
        }
    )


def _page_of(element) -> int | None:
    prov = getattr(element, "prov", None)
    if not prov:
        return None
    try:
        first = prov[0]
    except (IndexError, TypeError):
        return None
    page = getattr(first, "page_no", None)
    return int(page) if page is not None else None

def _merge_text_units(units: list[ParsedUnit]) -> list[ParsedUnit]:
    """Docling emits one element per paragraph, heading, or list item.
    Fed directly to the chunker, each becomes a tiny chunk (~15 tokens)
    with a weak embedding. Merge consecutive units that share the same
    section_title into a single ParsedUnit; the recursive chunker then
    splits by paragraph within that unit as normal.

    Section boundaries are respected: a new section_title flushes the
    buffer and starts a fresh unit. This keeps the heading as the first
    line of its section's content rather than mixing it with unrelated
    preceding text.
    """
    if not units:
        return []

    merged: list[ParsedUnit] = []
    buffer_text: list[str] = []
    current_section: str | None = None
    current_page: int | None = None

    def flush():
        nonlocal buffer_text, current_page
        if buffer_text:
            merged.append(ParsedUnit(
                text="\n\n".join(buffer_text),
                page_number=current_page,
                section_title=current_section,
            ))
            buffer_text = []
            current_page = None

    for unit in units:
        # Section change → flush and start new
        if unit.section_title != current_section:
            flush()
            current_section = unit.section_title

        # First unit of a new section: use its page number as the section's
        if current_page is None:
            current_page = unit.page_number

        buffer_text.append(unit.text)

    flush()
    return merged

def _safe_caption(element) -> str | None:
    """Docling exposes `caption_text` as a method, not a property. Call it,
    coerce to string, and swallow failures — captions are a nice-to-have,
    not a parse requirement."""
    raw = getattr(element, "caption_text", None)
    if raw is None:
        return None
    try:
        text = raw() if callable(raw) else raw
    except Exception:
        return None
    if not text:
        return None
    return str(text).strip() or None

def _convert_with_docling(content: bytes, suffix: str, document_id: str):
    """Docling infers format from file extension, so materialize bytes to
    a temp file with the right suffix."""
    converter = _get_docling_converter()
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        return converter.convert(tmp_path)
    except Exception as exc:
        raise IngestionError(document_id, f"Docling parse failed: {exc}") from exc
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def _iter_docling_items(doc):
    """Some Docling versions yield (element, level) tuples; others bare
    elements. Normalize to bare elements."""
    for item in doc.iterate_items():
        yield item[0] if isinstance(item, tuple) else item


def _walk_docling(doc, document_id: str) -> ParsedDocument:
    out = ParsedDocument()
    current_section: str | None = None

    for element in _iter_docling_items(doc):
        cls = type(element).__name__

        if cls == "SectionHeaderItem":
            current_section = (getattr(element, "text", "") or "").strip() or current_section
            if current_section:
                out.text_units.append(ParsedUnit(
                    text=current_section,
                    page_number=_page_of(element),
                    section_title=current_section,
                ))
            continue

        if cls == "TextItem":
            text = (getattr(element, "text", "") or "").strip()
            if text:
                out.text_units.append(ParsedUnit(
                    text=text,
                    page_number=_page_of(element),
                    section_title=current_section,
                ))
            continue

        if cls == "ListItem":
            text = (getattr(element, "text", "") or "").strip()
            if text:
                out.text_units.append(ParsedUnit(
                    text=f"- {text}",
                    page_number=_page_of(element),
                    section_title=current_section,
                ))
            continue

        if cls == "TableItem":
            try:
                md = (element.export_to_markdown(doc) or "").strip()
            except Exception as exc:
                logger.warning("table_markdown_failed",
                               extra={"document_id": document_id, "error": str(exc)})
                continue
            if md:
                out.table_units.append(TableUnit(
                    markdown=md,
                    page_number=_page_of(element),
                    section_title=current_section,
                ))
            continue

        if cls == "PictureItem":
            try:
                img = element.get_image(doc)
            except Exception as exc:
                logger.warning("picture_extract_failed",
                               extra={"document_id": document_id, "error": str(exc)})
                continue
            if img is None:
                continue
            buf = BytesIO()
            img.save(buf, format="PNG")
            out.picture_units.append(PictureUnit(
                image_bytes=buf.getvalue(),
                page_number=_page_of(element),
                section_title=current_section,
                caption=_safe_caption(element),
            ))
            continue

        # Other element types (CodeItem, FootnoteItem, FormulaItem)
        # ignored for now.

    out.text_units = _merge_text_units(out.text_units)
    return out



# ------------------------------------------------------------------ public

def parse_pdf(content: bytes, document_id: str) -> ParsedDocument:
    result = _convert_with_docling(content, ".pdf", document_id)
    parsed = _walk_docling(result.document, document_id)
    if not (parsed.text_units or parsed.table_units or parsed.picture_units):
        raise IngestionError(
            document_id,
            "No extractable content found in PDF (may be scanned/image-only).",
        )
    return parsed


def parse_docx(content: bytes, document_id: str) -> ParsedDocument:
    result = _convert_with_docling(content, ".docx", document_id)
    parsed = _walk_docling(result.document, document_id)
    if not (parsed.text_units or parsed.table_units or parsed.picture_units):
        raise IngestionError(document_id, "No extractable content found in DOCX.")
    return parsed


def parse_text(content: bytes, document_id: str) -> ParsedDocument:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise IngestionError(document_id, f"File is not valid UTF-8 text: {exc}") from exc
    if not text.strip():
        raise IngestionError(document_id, "File is empty.")
    return ParsedDocument(text_units=[ParsedUnit(text=text)])



# Namespace for the WordprocessingDrawing docPr element (alt-text container)
_WP_DOCPR_NS = (
    "{http://schemas.openxmlformats.org/drawingml/2006/"
    "wordprocessingDrawing}docPr"
)


def extract_docx_alt_texts(content: bytes) -> list[str]:
    """python-docx supplementary extractor for image alt text.

    Docling does not populate PictureItem.caption_text for DOCX. This
    walks the document XML directly for wp:docPr elements, which is the
    DrawingML container for image descriptions — works for both inline
    and floating images (inline_shapes only sees inline ones).

    Returns alt texts in document order; '' for images without alt text.
    """
    try:
        from docx import Document as DocxDocument
        from io import BytesIO as _BytesIO

        d = DocxDocument(_BytesIO(content))
        alts: list[str] = []
        for docPr in d.element.body.iter(_WP_DOCPR_NS):
            descr = (docPr.get("descr") or "").strip()
            alts.append(descr)
        return alts
    except Exception as exc:
        logger.warning("docx_alt_text_extract_failed", extra={"error": str(exc)})
        return []

PARSERS = {
    "application/pdf": parse_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": parse_docx,
    "text/plain": parse_text,
    "text/markdown": parse_text,
}


def parse_document(content_type: str, content: bytes, document_id: str) -> ParsedDocument:
    parser = PARSERS.get(content_type)
    if parser is None:
        raise IngestionError(
            document_id, f"No parser registered for content type '{content_type}'."
        )
    return parser(content, document_id)